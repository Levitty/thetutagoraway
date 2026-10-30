"""
Step 1 — read every document and pull out header fields and line items.

These are hand-typed Word and Excel files, so nothing is guaranteed: the word
"Quantity" may be "QTY", "Qty.", or a column with no header at all; a date may
be "12/03/2026", "12 March 2026" or "12.3.26"; the customer may be on a line
beginning "To:", "M/S", "CUSTOMER:" or simply sitting alone under the letterhead.

The rule throughout: when a field cannot be read with confidence, record the
document as unreadable WITH A REASON rather than guess. A guessed date silently
corrupts every gap calculation downstream, and nobody would ever find it.

Nothing here is specific to one customer or one item. All naming is handled in
clean.py so the mapping can be reviewed and corrected in one place.
"""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field, asdict
from datetime import datetime, date
from pathlib import Path

# ----------------------------------------------------------------- doc types
# Matched against the FILE NAME, longest/most specific first. Order matters:
# "CASH SALE INV 12" must not be read as an invoice.
DOC_PATTERNS = [
    ("cash_sale", re.compile(r"\bcash[\s_-]*sale", re.I)),
    ("statement", re.compile(r"\bstatement", re.I)),
    ("quote", re.compile(r"\b(quote|quotation|qtn|pro[\s_-]*forma)", re.I)),
    ("delivery_note", re.compile(r"\b(delivery[\s_-]*note|d\.?\s?note|\bdn\b|\bdl\b)", re.I)),
    ("invoice", re.compile(r"\b(invoice|invo|inv)\b|\binv\d", re.I)),
]

# Documents that represent a real movement of goods, i.e. an order.
ORDER_TYPES = {"invoice", "delivery_note", "cash_sale"}
# Kept, but never counted as orders.
NON_ORDER_TYPES = {"quote", "statement"}


def classify(name: str) -> str:
    for kind, pat in DOC_PATTERNS:
        if pat.search(name):
            return kind
    return "unknown"


# ---------------------------------------------------------------------- dates
_MONTHS = ("jan feb mar apr may jun jul aug sep oct nov dec").split()

DATE_PATTERNS = [
    # 12/03/2026, 12-3-26, 12.03.2026  -> DAY FIRST (Kenyan convention)
    (re.compile(r"\b(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})\b"), "dmy"),
    # 2026-03-12
    (re.compile(r"\b(\d{4})[/.\-](\d{1,2})[/.\-](\d{1,2})\b"), "ymd"),
    # 12 March 2026 / 12th Mar 26
    (re.compile(r"\b(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]{3,9})\.?,?\s+(\d{2,4})\b"), "dMy"),
    # March 12, 2026
    (re.compile(r"\b([A-Za-z]{3,9})\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{2,4})\b"), "Mdy"),
]


def _year(y: int) -> int:
    # Two-digit years in this data are 20xx. "26" is 2026, not 1926.
    return y + 2000 if y < 100 else y


def parse_date(text: str) -> date | None:
    """First plausible date in the text, or None. Day-first by convention.

    Day-first matters: 03/04/2026 is 3 April here, not 4 March. Getting it
    wrong shifts a gap by a month and the error is invisible in the output.
    """
    if not text:
        return None
    for pat, order in DATE_PATTERNS:
        for m in pat.finditer(text):
            try:
                if order == "dmy":
                    d, mo, y = int(m[1]), int(m[2]), _year(int(m[3]))
                    # A "day" above 12 in the second slot means it was written
                    # month-first after all; swap rather than throw it away.
                    if d > 31 or mo > 12:
                        if mo <= 31 and d <= 12:
                            d, mo = mo, d
                        else:
                            continue
                elif order == "ymd":
                    y, mo, d = _year(int(m[1])), int(m[2]), int(m[3])
                elif order == "dMy":
                    d, y = int(m[1]), _year(int(m[3]))
                    mo = _MONTHS.index(m[2][:3].lower()) + 1
                else:  # Mdy
                    mo = _MONTHS.index(m[1][:3].lower()) + 1
                    d, y = int(m[2]), _year(int(m[3]))
                dt = date(y, mo, d)
            except (ValueError, IndexError):
                continue
            # A document dated outside the plausible trading window is almost
            # always a misparse (a phone number, an amount, a reference).
            if date(2018, 1, 1) <= dt <= date(2030, 12, 31):
                return dt
    return None


# ------------------------------------------------------------------- numbers
def parse_num(raw) -> float | None:
    """'1,240.00', 'KES 1,240/=', '(50)' -> number. None if not a number."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return float(raw)
    s = str(raw).strip()
    if not s:
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = re.sub(r"[^0-9.\-]", "", s)
    if s in ("", "-", ".", "-."):
        return None
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if neg else v


# -------------------------------------------------------------- header fields
DOC_NO_PATTERNS = [
    re.compile(r"(?:invoice|inv|cash\s*sale|delivery\s*note|d\.?note|quote|quotation|statement)"
               r"\s*(?:no\.?|number|#|:)\s*[:#]?\s*([A-Za-z0-9/\-]{2,20})", re.I),
    re.compile(r"\b(?:no\.?|number|#)\s*[:#]?\s*([A-Za-z0-9/\-]{3,20})", re.I),
]

CUSTOMER_PATTERNS = [
    re.compile(r"^\s*(?:to|m/?s|messrs|customer|client|bill\s*to|sold\s*to)\s*[:.\-]\s*(.+)$", re.I),
]

# Lines that are the supplier's own letterhead, never the customer.
SELF_WORDS = re.compile(r"rycedental|ryce\s*dental|p\.?o\.?\s*box|nairobi\s*,?\s*kenya|"
                        r"tel\b|email|www\.|@|pin\s*no|vat\s*no", re.I)

# Column headers we expect in a line-item table.
COL_HINTS = {
    "description": ("description", "item", "particulars", "details", "product", "goods"),
    "qty": ("qty", "quantity", "units", "no.", "nos", "pcs"),
    "unit_price": ("unit price", "unitprice", "u/price", "rate", "price", "u.p", "unit cost"),
    "total": ("total", "amount", "value", "line total", "sub total", "subtotal"),
}
# Rows that close a table rather than describe goods.
TOTAL_ROW = re.compile(r"^\s*(sub\s*total|total|vat|v\.a\.t|grand\s*total|discount|amount\s*due|"
                       r"balance|less|add)\b", re.I)


@dataclass
class LineItem:
    description: str
    qty: float | None = None
    unit_price: float | None = None
    total: float | None = None


@dataclass
class Document:
    source_file: str
    doc_type: str
    doc_no: str | None = None
    doc_date: date | None = None
    customer: str | None = None
    items: list[LineItem] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def as_rows(self) -> list[dict]:
        out = []
        for it in self.items:
            out.append({
                "date": self.doc_date, "doc_type": self.doc_type, "doc_no": self.doc_no,
                "customer": self.customer, "description": it.description, "qty": it.qty,
                "unit_price": it.unit_price, "total": it.total, "source_file": self.source_file,
            })
        return out


# ------------------------------------------------------------------ docx read
def _docx_text_and_tables(path: Path):
    """Paragraphs (body + headers + footers) and tables, as plain structures.

    Headers and footers matter: on several of these templates the invoice
    number and date live in the header, so a body-only read loses both and the
    document looks unreadable when it is not.
    """
    import docx  # python-docx
    d = docx.Document(str(path))

    paras = [p.text for p in d.paragraphs]
    for section in d.sections:
        for part in (section.header, section.footer,
                     getattr(section, "first_page_header", None),
                     getattr(section, "first_page_footer", None),
                     getattr(section, "even_page_header", None),
                     getattr(section, "even_page_footer", None)):
            if part is None:
                continue
            try:
                paras += [p.text for p in part.paragraphs]
                for t in part.tables:
                    for r in t.rows:
                        paras += [c.text for c in r.cells]
            except Exception:
                pass

    tables = []
    for t in d.tables:
        rows = []
        for r in t.rows:
            # A merged cell repeats its text across the span. Collapsing runs
            # of identical adjacent cells keeps a 5-column table 5 columns wide
            # instead of silently becoming 8.
            cells, last = [], object()
            for c in r.cells:
                txt = c.text.strip()
                if txt and txt == last:
                    continue
                cells.append(txt)
                last = txt
            rows.append(cells)
        tables.append(rows)
    return paras, tables


# ----------------------------------------------------------------- excel read
def _excel_rows(path: Path) -> list[list]:
    rows = []
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                rows.append(["" if v is None else v for v in r])
            rows.append([])            # blank line between sheets
        wb.close()
    else:                              # .xls
        import xlrd
        book = xlrd.open_workbook(str(path))
        for sh in book.sheets():
            for i in range(sh.nrows):
                rows.append([c.value for c in sh.row(i)])
            rows.append([])
    return rows


# ------------------------------------------------------------- table reading
def _header_map(cells: list[str]) -> dict[str, int] | None:
    """If this row looks like a column header, say which column is which."""
    low = [str(c).strip().lower() for c in cells]
    found = {}
    for field_name, hints in COL_HINTS.items():
        for i, c in enumerate(low):
            if not c:
                continue
            if any(c == h or c.startswith(h) or h in c for h in hints):
                found.setdefault(field_name, i)
                break
    # A description column plus at least one number column is enough to trust.
    if "description" in found and ({"qty", "unit_price", "total"} & found.keys()):
        return found
    return None


def _items_from_rows(rows: list[list]) -> tuple[list[LineItem], list[str]]:
    items, notes = [], []
    hdr, hdr_at = None, -1
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in r]
        if not any(cells):
            continue
        if hdr is None:
            m = _header_map(cells)
            if m:
                hdr, hdr_at = m, i
            continue
        joined = " ".join(cells)
        if TOTAL_ROW.match(joined.strip()):
            continue                                  # a totals row, not goods
        def at(key):
            j = hdr.get(key)
            return cells[j] if j is not None and j < len(cells) else ""
        desc = at("description").strip()
        if not desc or len(desc) < 2:
            continue
        if desc.lower() in COL_HINTS["description"]:   # a repeated header row
            continue
        qty = parse_num(at("qty"))
        up = parse_num(at("unit_price"))
        tot = parse_num(at("total"))
        # A line with a description but no numbers at all is usually a note or
        # a continuation of the description above it.
        if qty is None and up is None and tot is None:
            notes.append(f"line with no figures ignored: {desc[:60]!r}")
            continue
        # Fill in the third figure when two are present and consistent.
        if tot is None and qty is not None and up is not None:
            tot = round(qty * up, 2)
        elif up is None and qty not in (None, 0) and tot is not None:
            up = round(tot / qty, 2)
        elif qty is None and up not in (None, 0) and tot is not None:
            q = tot / up
            qty = round(q) if abs(q - round(q)) < 0.01 else round(q, 2)
        items.append(LineItem(desc, qty, up, tot))
    if hdr is None:
        notes.append("no line-item table header found")
    return items, notes


def _customer_from(paras: list[str]) -> str | None:
    for line in paras:
        s = (line or "").strip()
        if not s:
            continue
        for pat in CUSTOMER_PATTERNS:
            m = pat.match(s)
            if m:
                cand = m.group(1).strip(" .:-")
                if cand and not SELF_WORDS.search(cand) and len(cand) > 2:
                    return cand
    return None


def _doc_no_from(paras: list[str], fallback_name: str) -> str | None:
    blob = "\n".join(p for p in paras if p)
    for pat in DOC_NO_PATTERNS:
        m = pat.search(blob)
        if m:
            return m.group(1).strip(" .:-")
    # Fall back to digits in the file name, which is how most of these are
    # actually identified in practice.
    m = re.search(r"(\d{2,6})", fallback_name)
    return m.group(1) if m else None


def read_document(path: Path) -> Document:
    doc = Document(source_file=path.name, doc_type=classify(path.name))
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            paras, tables = _docx_text_and_tables(path)
            flat_rows = [r for t in tables for r in t]
        elif suffix in (".xlsx", ".xls"):
            rows = _excel_rows(path)
            flat_rows = rows
            paras = [" ".join(str(c) for c in r if c not in (None, "")) for r in rows]
        else:
            doc.notes.append(f"unsupported file type {suffix}")
            return doc
    except (zipfile.BadZipFile, KeyError) as e:
        doc.notes.append(f"corrupt or password-protected: {e}")
        return doc
    except Exception as e:                                   # noqa: BLE001
        doc.notes.append(f"could not open: {type(e).__name__}: {e}")
        return doc

    blob = "\n".join(p for p in paras if p)
    doc.doc_date = parse_date(blob)
    doc.customer = _customer_from(paras)
    doc.doc_no = _doc_no_from(paras, path.stem)
    doc.items, notes = _items_from_rows(flat_rows)
    doc.notes += notes

    # A date is the one field nothing downstream can work without.
    if doc.doc_date is None:
        doc.notes.append("NO DATE FOUND")
    if not doc.customer:
        doc.notes.append("no customer line found")
    if not doc.items and doc.doc_type in ORDER_TYPES:
        doc.notes.append("no line items read")
    return doc


def walk(folder: Path) -> list[Document]:
    """Every readable document under `folder`, recursively."""
    docs = []
    for p in sorted(folder.rglob("*")):
        if p.is_dir() or p.name.startswith("~$") or p.name.startswith("."):
            continue                                   # ~$ = Word lock file
        if p.suffix.lower() in (".docx", ".xlsx", ".xls"):
            docs.append(read_document(p))
    return docs
